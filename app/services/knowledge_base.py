
from pathlib import Path
import hashlib
import re
import threading
from typing import Any

import chromadb
from docx import Document
from sentence_transformers import SentenceTransformer


# --------------------------------------------------
# 1. Paths and configuration
# --------------------------------------------------

# app/services/knowledge_base.py -> app/
APP_DIR = Path(__file__).resolve().parents[1]

KNOWLEDGE_DIR = APP_DIR / "knowledge"
CHROMA_DIR = APP_DIR / "chroma_db"

COLLECTION_NAME = "linguaai_knowledge"

# Multilingual embedding model.
# Supports semantic retrieval across many languages.
EMBEDDING_MODEL_NAME = (
    "sentence-transformers/"
    "paraphrase-multilingual-mpnet-base-v2"
)

# The embedding model truncates around 128 tokens, so keep chunks small.
CHUNK_SIZE = 70        # Approximate words per chunk
CHUNK_OVERLAP = 15     # Words shared between adjacent oversized chunks

# Bump to force re-indexing when indexing logic changes.
INDEX_VERSION = "2"

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}

_HEADER_RE = re.compile(
    r"^\s*(category|target[ _]locale)\s*:\s*(.+?)\s*$",
    re.IGNORECASE,
)

KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
CHROMA_DIR.mkdir(parents=True, exist_ok=True)


# --------------------------------------------------
# 2. Initialize persistent ChromaDB
# --------------------------------------------------

_client = chromadb.PersistentClient(
    path=str(CHROMA_DIR)
)

_collection = _client.get_or_create_collection(
    name=COLLECTION_NAME,
    metadata={"hnsw:space": "cosine"},
)


# --------------------------------------------------
# 3. Load the embedding model lazily
# --------------------------------------------------

_embedder = None
_knowledge_index_lock = threading.Lock()
_knowledge_index_ready = False


def get_embedder() -> SentenceTransformer:
    """Load the multilingual embedding model once."""
    global _embedder

    if _embedder is None:
        _embedder = SentenceTransformer(
            EMBEDDING_MODEL_NAME
        )

    return _embedder


# --------------------------------------------------
# 4. Extract text from supported files
# --------------------------------------------------

def extract_text(file_path: str | Path) -> str:
    """
    Extract text from TXT, Markdown, DOCX, and PDF files.

    Note: scanned PDFs require OCR, which is not implemented
    in this initial version.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"Knowledge file not found: {path}"
        )

    extension = path.suffix.lower()

    if extension in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig")

    if extension == ".docx":
        document = Document(str(path))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    if extension == ".pdf":
        try:
            import fitz
        except Exception as exc:
            raise RuntimeError(
                "PDF support is unavailable (PyMuPDF failed to load: "
                f"{exc}). Use TXT, MD or DOCX instead."
            ) from exc

        with fitz.open(str(path)) as pdf:
            return "\n".join(
                page.get_text("text")
                for page in pdf
            )

    raise ValueError(
        f"Unsupported file extension: {extension}. "
        f"Supported formats: {sorted(SUPPORTED_EXTENSIONS)}"
    )


# --------------------------------------------------
# 5. Split text into chunks
# --------------------------------------------------

def parse_header(text: str) -> tuple[dict[str, str], str]:
    """
    Extract 'Category:' and 'Target locale:' header lines.

    Returns (metadata, remaining_text).
    """
    meta: dict[str, str] = {}
    body_lines = []

    for line in text.splitlines():
        match = _HEADER_RE.match(line)

        if match:
            key = match.group(1).lower().replace(" ", "_")
            meta.setdefault(key, match.group(2))
        else:
            body_lines.append(line)

    return meta, "\n".join(body_lines)


def _split_words(
    text: str, chunk_size: int, overlap: int
) -> list[str]:
    words = re.findall(r"\S+", text)
    chunks = []
    start = 0

    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end]).strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(words):
            break

        start = end - overlap

    return chunks


def chunk_text(
    text: str,
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[str]:
    """
    Split text into chunks that respect blank-line-separated entries.

    Small entries are packed together up to chunk_size words, so a
    glossary entry is never cut in half. Entries longer than
    chunk_size are split by words with overlap.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive.")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be non-negative and smaller "
            "than chunk_size."
        )

    paragraphs = [
        p.strip()
        for p in re.split(r"\n\s*\n", text)
        if p.strip()
    ]

    chunks: list[str] = []
    current: list[str] = []
    current_words = 0

    def flush():
        nonlocal current, current_words
        if current:
            chunks.append("\n\n".join(current))
        current, current_words = [], 0

    for paragraph in paragraphs:
        size = len(paragraph.split())

        if size > chunk_size:
            flush()
            chunks.extend(
                _split_words(paragraph, chunk_size, overlap)
            )
            continue

        if current_words + size > chunk_size:
            flush()

        current.append(paragraph)
        current_words += size

    flush()

    return chunks


# --------------------------------------------------
# 6. Index one knowledge document
# --------------------------------------------------

def index_document(
    file_path: str | Path,
    category: str | None = None,
    source_language: str = "unknown",
    target_locale: str | None = None,
    force: bool = False,
) -> int:
    """
    Extract, chunk, embed, and store a knowledge document.

    Category and target locale are taken from explicit arguments,
    then from 'Category:' / 'Target locale:' header lines in the
    file, then default to 'general' / 'all'.

    Returns the number of indexed chunks. If the same version of
    the document is already indexed, it is skipped (unless force).
    Otherwise its previous chunks are replaced.
    """
    path = Path(file_path).resolve()

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported knowledge file: {path.suffix}"
        )

    text = extract_text(path)

    if not text.strip():
        raise ValueError(
            f"No extractable text found in {path.name}"
        )

    header, body = parse_header(text)
    category = category or header.get("category") or "general"
    target_locale = (
        target_locale or header.get("target_locale") or "all"
    )

    chunks = chunk_text(body)

    if not chunks:
        raise ValueError(
            f"No indexable content found in {path.name}"
        )

    # Identify this version of the document and its indexing settings.
    file_hash = hashlib.sha256(
        path.read_bytes()
        + f"|{INDEX_VERSION}|{category}|{target_locale}"
          f"|{source_language}".encode("utf-8")
    ).hexdigest()

    # Keep the collection's source identifier stable.
    source_id = path.name

    if not force:
        existing = _collection.get(
            where={"source": source_id},
            include=["metadatas"],
        )
        existing_metas = existing.get("metadatas") or []

        if (
            existing_metas
            and all(
                m.get("file_hash") == file_hash
                for m in existing_metas
            )
            and len(existing_metas) == len(chunks)
        ):
            return len(chunks)

    ids = [
        f"{file_hash[:20]}_{index}"
        for index in range(len(chunks))
    ]

    embeddings = get_embedder().encode(
        chunks,
        normalize_embeddings=True,
        show_progress_bar=False,
    ).tolist()

    metadatas: list[dict[str, Any]] = [
        {
            "source": source_id,
            "file_hash": file_hash,
            "category": category,
            "source_language": source_language,
            "target_locale": target_locale,
            "chunk_index": index,
        }
        for index in range(len(chunks))
    ]

    # Remove old chunks for this source before writing
    # the current version.
    _collection.delete(
        where={"source": source_id}
    )

    _collection.upsert(
        ids=ids,
        documents=chunks,
        embeddings=embeddings,
        metadatas=metadatas,
    )

    return len(chunks)


# --------------------------------------------------
# 7. Index all knowledge documents
# --------------------------------------------------

def index_knowledge_directory() -> dict[str, int]:
    """
    Index supported files inside app/knowledge/.

    Metadata comes from each file's header lines. Unchanged files
    are skipped, and chunks of files that no longer exist are removed.
    """
    summary: dict[str, int] = {}

    for path in sorted(KNOWLEDGE_DIR.iterdir()):
        if not path.is_file():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        summary[path.name] = index_document(path)

    for stale in set(list_indexed_sources()) - set(summary):
        _collection.delete(where={"source": stale})

    return summary


def ensure_knowledge_indexed() -> dict[str, int]:
    """Index knowledge on first RAG use instead of blocking API startup."""
    global _knowledge_index_ready
    if _knowledge_index_ready:
        return {}
    with _knowledge_index_lock:
        if _knowledge_index_ready:
            return {}
        summary = index_knowledge_directory()
        _knowledge_index_ready = True
        return summary


# --------------------------------------------------
# 8. Utility functions for retrieval and diagnostics
# --------------------------------------------------

def get_collection():
    """Return the ChromaDB collection for retrieval."""
    return _collection


def collection_count() -> int:
    """Return the number of indexed knowledge chunks."""
    return _collection.count()


def list_indexed_sources() -> list[str]:
    """Return the unique source filenames in the index."""
    if _collection.count() == 0:
        return []

    results = _collection.get(
        include=["metadatas"]
    )

    sources = {
        metadata["source"]
        for metadata in (results.get("metadatas") or [])
        if metadata and metadata.get("source")
    }

    return sorted(sources)
